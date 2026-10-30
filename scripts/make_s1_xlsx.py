#!/usr/bin/env python3
"""make_s1_xlsx.py - assemble S1 Table (Excel) from S1_Table_promoters.tsv and S1_Table_elements.tsv."""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

csv.field_size_limit(10 ** 8)
src, dst = Path(sys.argv[1]), Path(sys.argv[2])
P = list(csv.DictReader(open(src / "S1_Table_promoters.tsv", encoding="utf-8"), delimiter="\t"))
E = list(csv.DictReader(open(src / "S1_Table_elements.tsv", encoding="utf-8"), delimiter="\t"))

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2F5F9E")
LINK = Font(color="0563C1", underline="single")
MONO = Font(name="Consolas", size=10)
INT = {"start_site_GRCh38", "region_start_GRCh38", "region_end_GRCh38", "n_elements", "start_position", "end_position",
       "genomic_start_GRCh38", "genomic_end_GRCh38"} | {k for k in P[0] if k.startswith("has_")}
FLOAT = {"family_weight", "score", "threshold"}

wb = Workbook()
rd = wb.active
rd.title = "Read me"
readme = [
    ("S1 Table. Promoters of the main cohort, their sequences and every element called.", True),
    ("Marcet-Palacios M. Hierarchical Promoter Sequence Architecture Discovery (HPSAD). Code and data: "
     "https://github.com/mmarcet-palacios/hpsad, https://doi.org/10.5281/zenodo.23441932", False),
    ("", False),
    ("What is in this file", True),
    (f"Promoters sheet: one row per promoter ({len(P):,} promoters of the main cohort: full-length promoters of non-pseudogene "
     "genes whose start site is supported by EPDnew or FANTOM5 CAGE).", False),
    (f"Elements sheet: one row per element match ({len(E):,} matches). A promoter carries an element when at least one "
     "match lies in the element's positional zone; one element can have several overlapping matches.", False),
    ("The element calls are those of the paper (Fig 7B and S1 Fig); this table was regenerated from the archived code and "
     "checked to reproduce them exactly.", False),
    ("", False),
    ("Conventions", True),
    ("Genome: GRCh38 (UCSC hg38). Genomic coordinates are 1-based and inclusive.", False),
    ("Positions are relative to the start site: +1 is the start-site base, -1 the base just upstream; there is no position 0.", False),
    ("Sequences in the Promoters sheet are written 5' to 3' on the gene's strand. Lower-case bases are repeat-masked in the "
     "reference assembly; matches are never called on lower-case bases.", False),
    ("Elements sheet: matched_sequence_gene_strand is the match as read on the gene's strand; "
     "matched_sequence_genome_plus_strand is the same bases as shown by a genome browser (forward strand); "
     "matched_sequence_motif_orientation is the match read in the orientation of the motif consensus (differs only for "
     "antisense matches). context_gene_strand shows 10 bases on each side, with the match in brackets.", False),
    ("Motif consensus letters: IUPAC (R = A/G, Y = C/T, S = C/G, W = A/T, K = G/T, M = A/C, B/D/H/V = not A/C/G/T, N = any).", False),
    ("score / threshold: log-odds score of the match against the element's position weight matrix and the calling threshold; "
     "Initiator (Inr) and TCT are exact matches to their published consensus.", False),
    ("family_weight: 1 for a promoter with no duplicated copy; 1/n for each of n promoters that are copies of one another "
     "(duplicated gene families), so that each family counts once.", False),
    ("support: 6 of 6 = the element passed in all six independent discovery replicates; 4 of 6 = in four; start-site run = "
     "found in the start-site analysis (Fig 6).", False),
    ("", False),
    ("How to check any call by hand (under a minute)", True),
    ("1. In the Elements sheet, find the row (filter by gene_symbol and element).", False),
    ("2. Click the link in ucsc_element: the UCSC Genome Browser opens on the match with 20 bases on each side.", False),
    ("3. In UCSC choose View > DNA, then Get DNA: the bases at genomic_start..genomic_end are matched_sequence_genome_plus_strand.", False),
    ("4. Or search the promoter's sequence in the Promoters sheet (Ctrl+F) for the context shown in context_gene_strand.", False),
    ("", False),
    ("Worked examples", True),
    *[(f"{g} ({e['gene_strand']} strand): {e['element']} match {e['matched_sequence_gene_strand']} at {e['start_position']} to "
       f"{e['end_position']}, {e['chromosome']}:{int(e['genomic_start_GRCh38']):,}-{int(e['genomic_end_GRCh38']):,}; the browser "
       f"shows {e['matched_sequence_genome_plus_strand']} on the forward strand.", False)
      for g, e in [(g, next(x for x in E if x['gene_symbol'] == g and x['element'] == 'TATA')) for g in ("MYC", "HBB")]],
    ("", False),
    ("Links: ucsc_promoter opens the whole promoter region (-2000 to +100); ncbi_gene searches NCBI Gene for the gene symbol.", False),
]
for r, (txt, bold) in enumerate(readme, 1):
    c = rd.cell(row=r, column=1, value=txt)
    c.font = Font(bold=bold, size=12 if bold else 11)
    c.alignment = Alignment(wrap_text=True, vertical="top")
rd.column_dimensions["A"].width = 140


def sheet(title, rows, links, mono, widths):
    ws = wb.create_sheet(title)
    keys = list(rows[0])
    ws.append(keys)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for row in rows:
        vals = []
        for k in keys:
            v = row[k]
            if k in INT and v not in ("", None):
                v = int(v)
            elif k in FLOAT and v not in ("", None):
                try:
                    v = float(v)
                except ValueError:
                    pass
            vals.append(v)
        ws.append(vals)
    for k in links:
        j = keys.index(k) + 1
        for i in range(2, ws.max_row + 1):
            c = ws.cell(row=i, column=j)
            url = c.value
            c.value = "UCSC" if k.startswith("ucsc") else "NCBI"
            c.hyperlink = url
            c.font = LINK
    for k in mono:
        j = keys.index(k) + 1
        for i in range(2, ws.max_row + 1):
            ws.cell(row=i, column=j).font = MONO
    for j, k in enumerate(keys, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(k, max(10, min(24, len(k) + 2)))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


# put the links next to the identifiers so they are visible without scrolling
pk = list(P[0])
porder = ["promoter_id", "gene_symbol", "ucsc_promoter", "ncbi_gene"] + [k for k in pk if k not in
                                                                        ("promoter_id", "gene_symbol", "ucsc_promoter", "ncbi_gene")]
P = [{k: r[k] for k in porder} for r in P]
ek = list(E[0])
eorder = ["promoter_id", "gene_symbol", "element", "ucsc_element"] + [k for k in ek if k not in
                                                                     ("promoter_id", "gene_symbol", "element", "ucsc_element")]
E = [{k: r[k] for k in eorder} for r in E]
sheet("Promoters", P, ["ucsc_promoter", "ncbi_gene"],
      ["sequence_core_m300_p100", "sequence_upstream_m2000_m1", "sequence_downstream_p1_p100"],
      {"sequence_core_m300_p100": 60, "sequence_upstream_m2000_m1": 60, "sequence_downstream_p1_p100": 40,
       "architecture": 30, "promoter_id": 26, "biotype": 16})
sheet("Elements", E, ["ucsc_element"],
      ["motif_consensus", "matched_sequence_gene_strand", "matched_sequence_motif_orientation",
       "matched_sequence_genome_plus_strand", "context_gene_strand"],
      {"context_gene_strand": 40, "promoter_id": 26, "motif_consensus": 20, "matched_sequence_gene_strand": 20,
       "matched_sequence_motif_orientation": 20, "matched_sequence_genome_plus_strand": 20})
wb.save(dst)
print("saved", dst)
